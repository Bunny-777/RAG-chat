import io
import uuid
from typing import List, Tuple, Dict, Any, Optional
from langchain_core.documents import Document

from backend.app.core.exceptions import DocumentIngestionError
from backend.app.core.logging import get_logger

logger = get_logger(__name__)


def extract_text_from_pdf(content: bytes) -> Tuple[str, List[Document]]:
    """Extracts text page by page from PDF binary content using pypdf."""
    try:
        import pypdf
        reader = pypdf.PdfReader(io.BytesIO(content))
        docs = []
        full_text_list = []
        for i, page in enumerate(reader.pages):
            page_text = page.extract_text() or ""
            if page_text.strip():
                full_text_list.append(page_text)
                docs.append(
                    Document(
                        page_content=page_text,
                        metadata={"page": i + 1, "total_pages": len(reader.pages)},
                    )
                )
        if not docs:
            # Empty or scanned PDF fallback
            raise DocumentIngestionError("No readable text found in PDF. Scanned or image-only PDFs are not supported.")
        return "\n\n".join(full_text_list), docs
    except Exception as exc:
        if isinstance(exc, DocumentIngestionError):
            raise exc
        logger.error(f"Error parsing PDF file: {exc}")
        raise DocumentIngestionError(f"Failed to parse PDF document: {str(exc)}")


def extract_text_from_docx(content: bytes) -> Tuple[str, List[Document]]:
    """Extracts text from Word .docx file binary content."""
    try:
        import docx
        doc = docx.Document(io.BytesIO(content))
        paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
        full_text = "\n\n".join(paragraphs)
        if not full_text.strip():
            raise DocumentIngestionError("Word document contains no extractable text.")
        docs = [Document(page_content=full_text, metadata={"section": "full_body"})]
        return full_text, docs
    except Exception as exc:
        if isinstance(exc, DocumentIngestionError):
            raise exc
        logger.error(f"Error parsing DOCX file: {exc}")
        raise DocumentIngestionError(f"Failed to parse DOCX document: {str(exc)}")


def extract_text_from_plain(content: bytes, filename: str) -> Tuple[str, List[Document]]:
    """Extracts text from UTF-8 plain text, markdown, json, or csv file."""
    try:
        text = content.decode("utf-8", errors="replace")
        if not text.strip():
            raise DocumentIngestionError(f"File '{filename}' is empty.")
        docs = [Document(page_content=text, metadata={"filename": filename})]
        return text, docs
    except Exception as exc:
        if isinstance(exc, DocumentIngestionError):
            raise exc
        logger.error(f"Error reading plain text file {filename}: {exc}")
        raise DocumentIngestionError(f"Failed to read text file: {str(exc)}")


def create_document_documents(
    content: bytes,
    filename: str,
    source_id: Optional[str] = None,
) -> Tuple[List[Document], Dict[str, Any]]:
    """
    Parses document binary content according to file extension and attaches metadata.
    Returns (List[Document], metadata_dict).
    """
    sid = source_id or f"doc_{uuid.uuid4().hex[:8]}"
    lower_name = filename.lower()

    if lower_name.endswith(".pdf"):
        file_type = "pdf"
        full_text, docs = extract_text_from_pdf(content)
    elif lower_name.endswith(".docx") or lower_name.endswith(".doc"):
        file_type = "docx"
        full_text, docs = extract_text_from_docx(content)
    elif lower_name.endswith((".txt", ".md", ".markdown", ".json", ".csv", ".py", ".log")):
        file_type = "text"
        full_text, docs = extract_text_from_plain(content, filename)
    else:
        # Attempt fallback to text
        file_type = "text"
        full_text, docs = extract_text_from_plain(content, filename)

    # Attach normalized metadata to all documents
    for d in docs:
        d.metadata.update(
            {
                "source_id": sid,
                "source_type": "document",
                "filename": filename,
                "file_type": file_type,
            }
        )

    metadata = {
        "source_id": sid,
        "source_type": "document",
        "filename": filename,
        "file_type": file_type,
        "char_count": len(full_text),
        "doc_count": len(docs),
    }

    return docs, metadata
